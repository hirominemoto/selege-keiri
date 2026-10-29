"""
確認済みの「取り込み確認_企業と契約.xlsx」から、Supabase に流し込む SQL を作る。

使い方:
  python3 tools/build_seed.py <確認済み.xlsx> <出力.sql> [--email 社名,社名,...] [--on-demand 社名,社名,...]

  --email：いつもの送り方がメール（セレッジ担当が直接送る★の企業）の企業。正式名をカンマ区切りで
           実際の送り方は請求書ごとに選ぶので、ここで決めるのは「目安」
  --on-demand：毎月ではなく、請求するタイミングが来たときだけ請求する（都度請求）契約を持つ企業

出力した SQL は、Supabase の SQL Editor に貼って Run する（0001_init.sql の実行後）。
取引先の実データなので、出力はリポジトリに入れない（data/private/ に置く）。

・企業：「企業」シートのすべての行。取引の状態が「取引終了」なら取引終了として入れる
・旧Excelの表記ゆれ：正式名と違う表記を company_aliases に入れる
・企業のいつもの送り方（目安）：--email に挙げた企業（または「企業」シートの「送り方」列が「メール」の企業）はメール、ほかは郵送
・契約の請求のしかた：--on-demand に挙げた企業の契約（または「契約」シートの「請求のしかた」列が「都度」の契約）は都度請求、ほかは毎月
・契約：「契約」シートで「取り込む」が「はい」の行。全回数に達しているものは「終了」で入れる
・契約の「Excelで請求済みの最後の◯月分」：最後の請求日と最後の品名から決める
    品名に「◯月分」がある → その月（翌月分を前請求なら、請求日より後の月として年を決める）
    品名に「◯月分」がない → 最後の請求日の月
・同じ名前の企業・同じ企業の同じ契約（契約名と期）がすでにあれば入れない（何度実行しても重複しない）
"""

import datetime as dt
import re
import sys
import unicodedata

import openpyxl


def q(v):
    """SQL の文字列リテラル。空なら null"""
    if v is None:
        return "null"
    s = str(v).strip()
    if s == "":
        return "null"
    return "'" + s.replace("'", "''") + "'"


def num(v):
    if v is None or str(v).strip() == "":
        return "null"
    return str(int(round(float(unicodedata.normalize("NFKC", str(v)).replace(",", "")))))


def to_date(v):
    """セルの日付を date に。文字の日付（「2026/10/31」や打ち間違いの「202512/31」）も読む"""
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    digits = re.sub(r"\D", "", unicodedata.normalize("NFKC", str(v or "")))
    if len(digits) == 8:
        return dt.date(int(digits[:4]), int(digits[4:6]), int(digits[6:]))
    return None


def date_text(v):
    d = to_date(v)
    return f"{d:%Y/%m/%d}" if d else str(v or "")


def legacy_last_service_month(last_billed, desc, relation):
    """Excelで請求済みの最後の「◯月分」（その月の1日）。決められなければ None"""
    billed = to_date(last_billed)
    if billed is None:
        return None
    m = re.search(r"(\d{1,2})月分", unicodedata.normalize("NFKC", desc or ""))
    if m and "なし" not in (relation or ""):
        month = int(m.group(1))
        year = billed.year
        if "翌月" in (relation or "") and month < billed.month:
            year += 1          # 12月に1月分を請求したとき
        elif "翌月" not in (relation or "") and month > billed.month:
            year -= 1          # 念のため（請求月より先の月分は前の年とみなす）
        return dt.date(year, month, 1)
    return dt.date(billed.year, billed.month, 1)


def to_fullwidth_punct(s):
    """取り出すときに半角にした記号を、請求書らしい全角に戻す"""
    return s.replace("(", "（").replace(")", "）").replace(":", "：")


def make_template(desc, has_month):
    """最後の品名から、品名のひな形を作る"""
    t = unicodedata.normalize("NFKC", desc)
    t = re.sub(r"(\d+)回目\s*/\s*全\s*(\d+)\s*回", "{回数}回目/全{全回数}回", t)
    t = re.sub(r"(\d+)\s*(?:か月|ヶ月|ケ月)目", "{回数}か月目", t)
    if has_month:
        t = re.sub(r"(\d{1,2})月分", "{対象月}分", t)
    t = re.sub(r"(\d+)\s*名", "{人数}名", t, count=1)
    return to_fullwidth_punct(t.strip())


def main(src, out_path, email_companies=(), on_demand_companies=()):
    email_companies = {n.strip() for n in email_companies if n.strip()}
    on_demand_companies = {n.strip() for n in on_demand_companies if n.strip()}
    on_demand_found = set()
    wb = openpyxl.load_workbook(src, data_only=True)
    lines = [
        "-- セレッジ管理システム：旧Excelからの取り込み（企業・表記ゆれ・契約）",
        "-- 0001_init.sql を実行したあとに、SQL Editor で実行する",
        "begin;",
        "",
    ]

    # ---- 企業 ----
    ws = wb["企業"]
    head = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(head)}
    n_company = n_alias = n_email = 0
    email_found = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        name = row[idx["正式名"]]
        if not name:
            continue
        name = str(name).strip()
        active = (row[idx["取引の状態"]] or "取引中") != "取引終了"
        by_email = name in email_companies or (
            "送り方" in idx and "メール" in str(row[idx["送り方"]] or ""))
        if by_email:
            n_email += 1
            email_found.add(name)
        lines.append(
            "insert into public.companies (name, postal_code, address1, contact_name, send_method, is_active) values "
            f"({q(name)}, {q(row[idx['郵便番号']])}, {q(row[idx['住所']])}, {q(row[idx['部署・担当者']])}, "
            f"{q('email' if by_email else 'post')}, {'true' if active else 'false'}) on conflict (name) do nothing;"
        )
        n_company += 1
        variants = [v.strip() for v in str(row[idx["Excelでの表記"]] or "").split("、") if v.strip()]
        for v in dict.fromkeys(variants):
            if v != name:
                lines.append(
                    "insert into public.company_aliases (company_id, alias) "
                    f"select id, {q(v)} from public.companies where name = {q(name)} on conflict (alias) do nothing;"
                )
                n_alias += 1
    lines.append("")

    # ---- 契約 ----
    ws = wb["契約"]
    head = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(head)}
    n_contract = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[idx["企業"]] or (row[idx["取り込む"]] or "はい") != "はい":
            continue
        company = str(row[idx["企業"]]).strip()
        relation = row[idx["◯月分の関係"]] or ""
        offset = 1 if "翌月" in relation else 0
        has_month = "なし" not in relation
        last = row[idx["最後の回数"]]
        total = row[idx["全回数"]]
        desc = str(row[idx["最後の品名（元の表記）"]] or "")
        count_unit = "か月" if re.search(r"(か月|ヶ月|ケ月)目", desc) else "回"
        finished = last is not None and total is not None and int(last) >= int(total)
        cohort = row[idx["期"]]
        on_demand = company in on_demand_companies or (
            "請求のしかた" in idx and "都度" in str(row[idx["請求のしかた"]] or ""))
        if on_demand:
            on_demand_found.add(company)
        legacy = legacy_last_service_month(row[idx["最後の請求日"]], desc, relation)
        legacy_sql = f"'{legacy:%Y-%m-%d}'" if legacy else "null"
        lines.append(
            "insert into public.contracts (company_id, name, cohort, kind, quantity, unit_price, taxable, "
            "service_month_offset, total_count, count_unit, base_count, legacy_last_service_month, status, "
            "description_template, note) "
            f"select co.id, {q(row[idx['契約名']])}, {q(cohort)}, {q('on_demand' if on_demand else 'monthly')}, {num(row[idx['人数']]) if row[idx['人数']] else 1}, "
            f"{num(row[idx['単価（税抜）']])}, true, {offset}, {num(total)}, {q(count_unit)}, {num(last) if last else 0}, "
            f"{legacy_sql}, {q('ended' if finished else 'active')}, {q(make_template(desc, has_month))}, "
            f"{q('旧Excelから取り込み。最後の請求：' + date_text(row[idx['最後の請求日']]))} "
            f"from public.companies co where co.name = {q(company)} "
            "and not exists (select 1 from public.contracts c where c.company_id = co.id "
            f"and c.name = {q(row[idx['契約名']])} and coalesce(c.cohort, '') = coalesce({q(cohort)}, ''));"
        )
        n_contract += 1

    lines += ["", "commit;", ""]
    with open(out_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    print(f"企業 {n_company} 件（うち、いつもの送り方がメール {n_email} 件）、表記ゆれ {n_alias} 件、契約 {n_contract} 件の SQL を書き出しました → {out_path}")
    print(f"都度請求の契約がある企業：{'、'.join(sorted(on_demand_found)) or 'なし'}")
    missing = email_companies - email_found
    missing_od = on_demand_companies - on_demand_found
    if missing or missing_od:
        if missing:
            print("※ --email に挙げたが「企業」シートに見つからない企業：" + "、".join(sorted(missing)))
        if missing_od:
            print("※ --on-demand に挙げたが「契約」シートに見つからない企業：" + "、".join(sorted(missing_od)))
        sys.exit(2)


if __name__ == "__main__":
    args = sys.argv[1:]

    def take(flag):
        global args
        if flag not in args:
            return []
        i = args.index(flag)
        vals = args[i + 1].replace("、", ",").split(",") if i + 1 < len(args) else []
        args = args[:i] + args[i + 2:]
        return vals

    emails = take("--email")
    on_demand = take("--on-demand")
    if len(args) != 2:
        print(__doc__)
        sys.exit(1)
    main(args[0], args[1], emails, on_demand)
